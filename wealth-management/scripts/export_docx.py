#!/usr/bin/env python3
"""
Export a Strategy Request markdown file to a formatted Word document (.docx).

Usage:
    python3 export_docx.py input.md [output.docx]

If no output path is provided, the output file is named after the input file
with a .docx extension in the same directory.

Requires: python-docx (pip install python-docx)
"""

import sys
import re
import os
from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn


def parse_markdown_tables(text):
    """Parse markdown tables into list of rows (list of cells)."""
    lines = [l.strip() for l in text.strip().split("\n") if l.strip()]
    rows = []
    for line in lines:
        if line.startswith("|") and not re.match(r"^\|[\s\-:|]+\|$", line):
            cells = [c.strip() for c in line.split("|")[1:-1]]
            rows.append(cells)
    return rows


def set_cell_shading(cell, color_hex):
    """Set background shading on a table cell."""
    shading = cell._element.get_or_add_tcPr()
    shading_elem = shading.makeelement(
        qn("w:shd"),
        {
            qn("w:fill"): color_hex,
            qn("w:val"): "clear",
        },
    )
    shading.append(shading_elem)


def style_header_row(row, bg_color="1F3864"):
    """Style a table header row with dark background and white text."""
    for cell in row.cells:
        set_cell_shading(cell, bg_color)
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                run.font.bold = True
                run.font.size = Pt(9)


def add_table_from_rows(doc, headers, data_rows, first_col_width=None):
    """Add a formatted table to the document."""
    if not headers and not data_rows:
        return

    col_count = len(headers) if headers else len(data_rows[0]) if data_rows else 0
    if col_count == 0:
        return

    total_rows = (1 if headers else 0) + len(data_rows)
    table = doc.add_table(rows=total_rows, cols=col_count)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"

    row_idx = 0
    if headers:
        for j, header in enumerate(headers):
            cell = table.cell(0, j)
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(header)
            run.font.size = Pt(9)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        style_header_row(table.rows[0])
        row_idx = 1

    for data_row in data_rows:
        for j in range(min(len(data_row), col_count)):
            cell = table.cell(row_idx, j)
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(data_row[j])
            run.font.size = Pt(9)
            run.font.name = "Calibri"
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        row_idx += 1

    # Set font for header row
    if headers:
        for cell in table.rows[0].cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.name = "Calibri"

    return table


def process_section(doc, section_title, section_body):
    """Process a single section of the strategy request."""
    doc.add_heading(section_title, level=2)

    lines = section_body.strip().split("\n")

    # Separate table content from non-table text
    table_lines = []
    text_lines = []
    in_table = False

    for line in lines:
        stripped = line.strip()
        if stripped.startswith("|"):
            in_table = True
            table_lines.append(stripped)
        else:
            if in_table and table_lines:
                # Process accumulated table
                process_table_block(doc, table_lines)
                table_lines = []
                in_table = False
            if stripped:
                text_lines.append(stripped)
                p = doc.add_paragraph(stripped)
                p.style = doc.styles["Normal"]

    # Process any remaining table
    if table_lines:
        process_table_block(doc, table_lines)


def process_table_block(doc, table_lines):
    """Process a block of markdown table lines into a Word table."""
    # Parse header and separator
    all_rows = []
    separator_idx = None

    for i, line in enumerate(table_lines):
        if re.match(r"^\|[\s\-:|]+\|$", line):
            separator_idx = i
            continue
        cells = [c.strip() for c in line.split("|")[1:-1]]
        all_rows.append(cells)

    if not all_rows:
        return

    # If we found a separator, first row is headers
    if separator_idx is not None and separator_idx > 0:
        headers = all_rows[0]
        data_rows = all_rows[1:]
    else:
        headers = all_rows[0]
        data_rows = all_rows[1:]

    # Skip row-number first column if it only contains numbers
    has_row_numbers = True
    for row in data_rows:
        if row and not re.match(r"^\d*$", row[0].strip()):
            has_row_numbers = False
            break

    if has_row_numbers and len(headers) > 1:
        headers = headers[1:]
        data_rows = [row[1:] for row in data_rows if len(row) > 1]

    # Filter out empty rows
    data_rows = [row for row in data_rows if any(cell.strip() for cell in row)]

    add_table_from_rows(doc, headers, data_rows)
    doc.add_paragraph("")  # spacing after table


def convert_md_to_docx(md_path, docx_path):
    """Convert a strategy request markdown file to a Word document."""
    with open(md_path, "r") as f:
        content = f.read()

    doc = Document()

    # Set default font
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Calibri"
    font.size = Pt(10)

    # Set narrow margins
    for section in doc.sections:
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)

    # Style headings
    for level in range(1, 4):
        heading_style = doc.styles[f"Heading {level}"]
        heading_style.font.name = "Calibri"
        heading_style.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)

    doc.styles["Heading 1"].font.size = Pt(16)
    doc.styles["Heading 2"].font.size = Pt(12)

    # Extract title
    title_match = re.match(r"^#\s+(.+)$", content.strip(), re.MULTILINE)
    if title_match:
        title = title_match.group(1)
        doc.add_heading(title, level=1)
    else:
        doc.add_heading("SOA Summary Request", level=1)

    # Add date line
    from datetime import date

    p = doc.add_paragraph(f"Date: {date.today().strftime('%d %B %Y')}")
    p.runs[0].font.size = Pt(9)
    p.runs[0].font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    doc.add_paragraph("")  # spacer

    # Split into sections by ## headings
    sections = re.split(r"^##\s+", content, flags=re.MULTILINE)

    for section_text in sections[1:]:  # skip content before first ##
        lines = section_text.strip().split("\n", 1)
        section_title = lines[0].strip()
        section_body = lines[1] if len(lines) > 1 else ""

        process_section(doc, section_title, section_body)

    # Add footer
    footer_section = doc.sections[0]
    footer = footer_section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.text = "SOA Summary Request - Internal Working Paper - Not for Client Distribution"
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].font.size = Pt(7)
    p.runs[0].font.color.rgb = RGBColor(0x99, 0x99, 0x99)
    p.runs[0].font.name = "Calibri"

    doc.save(docx_path)
    return docx_path


def main():
    if len(sys.argv) < 2:
        print("Usage: python3 export_docx.py input.md [output.docx]")
        sys.exit(1)

    md_path = sys.argv[1]
    if not os.path.exists(md_path):
        print(f"Error: {md_path} not found")
        sys.exit(1)

    if len(sys.argv) >= 3:
        docx_path = sys.argv[2]
    else:
        docx_path = os.path.splitext(md_path)[0] + ".docx"

    output = convert_md_to_docx(md_path, docx_path)
    print(f"Exported: {output}")


if __name__ == "__main__":
    main()
